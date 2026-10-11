#!/usr/bin/env python3
"""Extract criterion-level GHSMTA scoring guidance from the approved DOCX."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document


CATEGORY_KEYS = {
    "MUSIC DIRECTION": "music_direction",
    "ORCHESTRA": "orchestra",
    "CHOREOGRAPHY": "choreography",
    "ENSEMBLE": "ensemble",
    "TECHNICAL EXECUTION": "technical_execution",
    "LIGHTING DESIGN": "lighting_design",
    "SOUND DESIGN": "sound_design",
    "SCENIC": "scenic_design",
    "COSTUMES": "costume_design",
    "LEADING ACTRESS": "leading_actress",
    "LEADING ACTOR": "leading_actor",
    "SUPPORTING PERFORMER A": "supporting_performer_a",
    "SUPPORTING PERFORMER B": "supporting_performer_b",
    "FEATURED PERFORMER": "featured_performer",
    "DIRECTION": "direction",
}


def normalized_text(value: str) -> str:
    return " ".join(value.split())


def cell_text(cell) -> str:
    return normalized_text(" ".join(paragraph.text for paragraph in cell.paragraphs))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--expected-criteria", type=int, default=58)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    document = Document(args.input)

    if len(document.tables) != 1 + (args.expected_criteria * 2):
        raise ValueError(
            "Expected one overview table and two tables per criterion; "
            f"found {len(document.tables)} tables."
        )

    score_labels: dict[int, str] = {}
    for row in document.tables[0].rows[1:]:
        score = int(cell_text(row.cells[0]))
        descriptor = cell_text(row.cells[1])
        score_labels[score] = re.split(r"\s+[—–-]\s+", descriptor, maxsplit=1)[0]

    criteria: list[dict[str, object]] = []
    current_category_key: str | None = None

    for paragraph in document.paragraphs:
        text = normalized_text(paragraph.text)
        if not text:
            continue

        if text in CATEGORY_KEYS:
            current_category_key = CATEGORY_KEYS[text]
            continue

        heading = re.match(r"^CRITERION (\d+) \| (.+)$", text)
        if heading:
            sort_order = int(heading.group(1))
            criterion_identity = (current_category_key, sort_order)
            previous_identity = (
                (criteria[-1]["category_key"], criteria[-1]["sort_order"])
                if criteria
                else None
            )
            if criterion_identity != previous_identity:
                if current_category_key is None:
                    raise ValueError(f"Criterion appears before a category: {text}")
                criteria.append(
                    {
                        "category_key": current_category_key,
                        "sort_order": sort_order,
                        "title": heading.group(2),
                        "assessment_focus": [],
                    }
                )
            continue

        if text.startswith("ASSESSMENT FOCUS "):
            if not criteria:
                raise ValueError("Assessment focus appears before a criterion.")
            criteria[-1]["assessment_focus"] = [
                item.strip()
                for item in text.removeprefix("ASSESSMENT FOCUS ").split(" • ")
                if item.strip()
            ]

    if len(criteria) != args.expected_criteria:
        raise ValueError(
            f"Expected {args.expected_criteria} criteria; found {len(criteria)}."
        )

    for index, criterion in enumerate(criteria):
        levels: list[dict[str, object]] = []
        for table in (
            document.tables[1 + (index * 2)],
            document.tables[2 + (index * 2)],
        ):
            if len(table.columns) != 3:
                raise ValueError(f"Criterion {index + 1} has an unexpected table shape.")
            for row in table.rows[1:]:
                score = int(cell_text(row.cells[0]))
                levels.append(
                    {
                        "score": score,
                        "label": score_labels[score],
                        "observable": cell_text(row.cells[1]),
                        "example": cell_text(row.cells[2]),
                    }
                )

        if [level["score"] for level in levels] != list(range(10, 0, -1)):
            raise ValueError(
                f"Criterion {index + 1} does not contain scores 10 through 1."
            )
        if len(criterion["assessment_focus"]) not in (4, 5):
            raise ValueError(
                f"Criterion {index + 1} has an unexpected assessment-focus count."
            )
        criterion["levels"] = levels

    payload = {
        "source_document": args.input.name,
        "season_year": "2026-2027",
        "criteria": criteria,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(
        f"Extracted {len(criteria)} criteria and "
        f"{sum(len(item['levels']) for item in criteria)} score levels."
    )


if __name__ == "__main__":
    main()
